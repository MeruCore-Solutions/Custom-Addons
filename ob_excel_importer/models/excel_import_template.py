import base64
import fnmatch
import ftplib
import io
import json
import logging
import posixpath

from datetime import date, datetime, time, timedelta

from dateutil import parser as date_parser
from dateutil.relativedelta import relativedelta
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

try:
    import paramiko
except ImportError:
    paramiko = None

from odoo import _, api, fields, models, Command
from odoo.exceptions import AccessDenied, AccessError, UserError, ValidationError
from odoo.tools.safe_eval import safe_eval

_logger = logging.getLogger(__name__)

SKIP = object()
TECHNICAL_FIELDS = {
    "id",
    "create_uid",
    "create_date",
    "write_uid",
    "write_date",
    "display_name",
    "__last_update",
}


class ImportRowError(Exception):
    def __init__(self, message, error_type="validation", mapped_vals=None, blocking=False):
        super().__init__(message)
        self.message = message
        self.error_type = error_type
        self.mapped_vals = mapped_vals or {}
        self.blocking = blocking


class ExcelImportTemplate(models.Model):
    _name = "ob.excel.import.template"
    _description = "Excel Import Template"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _order = "sequence, id"

    name = fields.Char(required=True, tracking=True)
    active = fields.Boolean(default=True)
    sequence = fields.Integer(default=10)
    model_id = fields.Many2one(
        "ir.model",
        string="Target Model",
        required=True,
        ondelete="cascade",
    )
    target_model = fields.Char(
        related="model_id.model",
        store=True,
        readonly=True,
    )
    sheet_name = fields.Char()
    header_row = fields.Integer(default=1, required=True)
    data_start_row = fields.Integer(default=2, required=True)
    batch_size = fields.Integer(default=500)
    import_mode = fields.Selection(
        [
            ("create", "Create Only"),
            ("update", "Update Only"),
            ("create_update", "Create or Update"),
        ],
        default="create_update",
        required=True,
    )
    duplicate_policy = fields.Selection(
        [
            ("skip", "Skip"),
            ("update", "Update"),
            ("error", "Error"),
        ],
        default="update",
        required=True,
    )
    external_key_column = fields.Char(
        help="Excel column/header used as unique source key.",
    )
    external_key_field = fields.Char(
        help="Target Odoo field used to find existing records.",
    )
    stop_on_first_error = fields.Boolean(
        default=False,
        help="Stop the whole import when the first blocking error is found.",
    )
    field_mapping_ids = fields.One2many(
        "ob.excel.import.field",
        "template_id",
        string="Field Mappings",
        copy=True,
    )
    log_ids = fields.One2many(
        "ob.excel.import.log",
        "template_id",
        string="Import Logs",
    )
    post_process_ids = fields.One2many(
        "ob.excel.import.post.process",
        "template_id",
        string="Post Processes",
    )
    record_mapping_ids = fields.One2many(
        "ob.excel.import.record.map",
        "template_id",
        string="Record Mappings",
    )
    responsible_user_id = fields.Many2one(
        "res.users",
        default=lambda self: self.env.user,
        required=True,
        tracking=True,
    )
    notify_on_error = fields.Boolean(default=True)
    create_activity_on_error = fields.Boolean(default=True)
    activity_type_id = fields.Many2one(
        "mail.activity.type",
        default=lambda self: self.env.ref("mail.mail_activity_data_todo", raise_if_not_found=False),
    )
    auto_import = fields.Boolean(default=False)
    interval_number = fields.Integer(default=1)
    interval_unit = fields.Selection(
        [
            ("minute", "Minute"),
            ("hour", "Hour"),
            ("day", "Day"),
            ("week", "Week"),
            ("month", "Month"),
        ],
        default="day",
        required=True,
    )
    next_execution_datetime = fields.Datetime()
    last_run_datetime = fields.Datetime(readonly=True)
    last_message = fields.Text(readonly=True)
    ftp_enabled = fields.Boolean(default=False)
    ftp_protocol = fields.Selection(
        [
            ("ftp", "FTP"),
            ("sftp", "SFTP"),
        ],
        default="ftp",
    )
    ftp_host = fields.Char()
    ftp_port = fields.Integer(default=21)
    ftp_username = fields.Char()
    ftp_password = fields.Char()
    ftp_path = fields.Char(default="/")
    ftp_filename_pattern = fields.Char(default="*.xlsx")
    ftp_archive_path = fields.Char()
    ftp_delete_after_import = fields.Boolean(default=False)
    ftp_passive_mode = fields.Boolean(default=True)
    pre_process_python = fields.Text(
        string="Before Mapping Python Code",
        default="""
# Executed before field mapping.
# Available variables:
# source_vals, env, template, datetime, date, time, relativedelta,
# UserError, Command, log, _logger
""".strip(),
    )
    row_filter_python = fields.Text(
        string="Row Filter Python Code",
        default="""
# Set skip_row = True to skip this row.
# Set skip_reason = "reason" to explain why.
""".strip(),
    )
    post_process_python = fields.Text(
        string="Before Create/Write Python Code",
        default="""
# Executed after field mapping and before create/write.
# Must mutate result_vals in place.
# Available variables:
# source_vals, result_vals, record, env, template, datetime, date, time,
# relativedelta, UserError, Command, log, _logger
""".strip(),
    )
    post_import_python = fields.Text(
        string="After Import Python Code",
        default="""
# Executed after the complete file import.
# Available variables:
# env, template, import_log, created_records, updated_records,
# failed_rows, datetime, date, time, relativedelta, UserError, Command,
# log, _logger
""".strip(),
    )
    log_count = fields.Integer(compute="_compute_counts")
    failed_line_count = fields.Integer(compute="_compute_counts")
    pending_post_process_count = fields.Integer(compute="_compute_counts")
    record_mapping_count = fields.Integer(compute="_compute_counts")

    @api.depends("log_ids", "post_process_ids", "record_mapping_ids")
    def _compute_counts(self):
        failed_line_model = self.env["ob.excel.import.failed.line"].sudo()
        for template in self:
            template.log_count = len(template.log_ids)
            template.pending_post_process_count = len(template.post_process_ids.filtered(lambda process: process.state == "pending"))
            template.record_mapping_count = len(template.record_mapping_ids)
            template.failed_line_count = failed_line_model.search_count([("template_id", "=", template.id)])

    @api.model_create_multi
    def create(self, vals_list):
        records = super().create(vals_list)
        for template in records:
            if template.auto_import and template.ftp_enabled and not template.next_execution_datetime:
                template.sudo()._update_next_execution()
        return records

    def write(self, vals):
        result = super().write(vals)
        schedule_fields = {"auto_import", "ftp_enabled", "interval_number", "interval_unit"}
        if schedule_fields & set(vals):
            for template in self:
                template.sudo()._update_next_execution()
        return result

    def action_open_import_wizard(self):
        self.ensure_one()
        return self._open_import_wizard()

    def action_open_test_import_wizard(self):
        self.ensure_one()
        return self._open_import_wizard(default_dry_run=True)

    def _open_import_wizard(self, default_dry_run=False):
        return {
            "type": "ir.actions.act_window",
            "name": _("Excel Import"),
            "res_model": "ob.excel.import.wizard",
            "view_mode": "form",
            "view_id": self.env.ref("ob_excel_importer.view_ob_excel_import_wizard_form").id,
            "target": "new",
            "context": {
                "default_template_id": self.id,
                "default_dry_run": default_dry_run,
            },
        }

    def action_view_logs(self):
        self.ensure_one()
        return {
            "type": "ir.actions.act_window",
            "name": _("Import Logs"),
            "res_model": "ob.excel.import.log",
            "view_mode": "list,form",
            "domain": [("template_id", "=", self.id)],
            "context": {"default_template_id": self.id},
        }

    def action_view_failed_lines(self):
        self.ensure_one()
        return {
            "type": "ir.actions.act_window",
            "name": _("Not Imported Lines"),
            "res_model": "ob.excel.import.failed.line",
            "view_mode": "list,form",
            "domain": [("template_id", "=", self.id)],
            "context": {"default_template_id": self.id},
        }

    def action_run_pending_post_processes(self):
        self.ensure_one()
        processes = self.post_process_ids.filtered(lambda process: process.state == "pending").sorted("priority")
        success_count = 0
        failed_count = 0
        for process in processes:
            try:
                process.execute()
                success_count += 1
            except Exception:
                failed_count += 1
        message = _("%s pending post process(es) executed, %s failed.") % (success_count, failed_count)
        return self._build_notification_action(
            title=_("Post Processes"),
            message=message,
            notif_type="warning" if failed_count else "success",
        )

    def action_generate_field_mappings(self):
        self.ensure_one()
        self._validate_template_model()
        existing_fields = set(self.field_mapping_ids.mapped("target_field"))
        target_model = self.env[self.target_model]
        values_list = []
        for field_name, field in target_model._fields.items():
            if field_name in TECHNICAL_FIELDS:
                continue
            if field_name in existing_fields:
                continue
            if not field.store or field.readonly:
                continue
            if field.type in ("binary", "reference", "many2one_reference", "properties"):
                continue
            values_list.append({
                "template_id": self.id,
                "excel_header": field_name,
                "target_field": field_name,
                "transform": self._guess_transform_for_field(field),
            })
        if values_list:
            self.env["ob.excel.import.field"].create(values_list)
        return self._build_notification_action(
            title=_("Field Mappings"),
            message=_("%s mapping(s) generated.") % len(values_list),
        )

    def run_import(self, file_content=None, filename=None, dry_run=False):
        self.ensure_one()
        log, summary = self._run_import_core(file_content=file_content, filename=filename, dry_run=dry_run)
        title = _("Test Import") if dry_run else _("Excel Import")
        notif_type = "success"
        if summary["state"] in ("failed",):
            notif_type = "danger"
        elif summary["state"] in ("partial",):
            notif_type = "warning"
        return self._build_notification_action(
            title=title,
            message=summary["message"],
            notif_type=notif_type,
            next_action=log and {
                "type": "ir.actions.act_window",
                "name": _("Import Log"),
                "res_model": "ob.excel.import.log",
                "res_id": log.id,
                "view_mode": "form",
                "target": "current",
            } or False,
        )

    def _run_import_core(self, file_content=None, filename=None, dry_run=False):
        self.ensure_one()
        filename = filename or _("import.xlsx")
        try:
            self._validate_template_configuration()
            binary_file = self._decode_binary_file(file_content)
        except Exception as exc:
            message = str(exc)
            log = self._create_system_failed_log(message, filename=filename)
            self._notify_import_blockage(log=log, message=message, exception=exc)
            return log, {
                "state": "failed",
                "message": message,
                "success_count": 0,
                "failed_count": 0,
                "skipped_count": 0,
            }

        log = self._create_import_log(filename=filename)
        target_model = self.env[self.target_model]
        created_records = target_model.browse()
        updated_records = target_model.browse()
        failed_rows = self.env["ob.excel.import.failed.line"].sudo()
        success_count = 0
        failed_count = 0
        skipped_count = 0
        total_rows = 0
        blocked_error = None
        workbook = None
        try:
            workbook = self._read_excel_file(binary_file, filename)
            sheet = self._get_workbook_sheet(workbook)
            headers = self._extract_headers(sheet)
            for source_vals in self._iter_excel_rows(sheet, headers):
                total_rows += 1
                row_number = source_vals.get("_row_number")
                try:
                    with self.env.cr.savepoint():
                        result = self._process_single_source_row(
                            source_vals=source_vals,
                            log=log,
                            row_number=row_number,
                            filename=filename,
                            dry_run=dry_run,
                        )
                except Exception as exc:
                    error = self._normalize_row_exception(exc)
                    failed_line = self._create_failed_line(
                        log=log,
                        row_number=row_number,
                        source_vals=source_vals,
                        mapped_vals=error.mapped_vals,
                        error_message=error.message,
                        error_type=error.error_type,
                        filename=filename,
                    )
                    self._create_log_line(
                        log=log,
                        row_number=row_number,
                        state="failed",
                        source_vals=source_vals,
                        mapped_vals=error.mapped_vals,
                        message=error.message,
                    )
                    failed_rows |= failed_line
                    failed_count += 1
                    if self.stop_on_first_error or error.blocking:
                        blocked_error = error
                        self._notify_import_blockage(
                            log=log,
                            message=error.message,
                            failed_line=failed_line,
                            exception=exc,
                        )
                        break
                    continue

                if result["status"] == "skipped":
                    skipped_count += 1
                else:
                    success_count += 1
                    if result["record"]:
                        if result["operation"] == "create":
                            created_records |= result["record"]
                        elif result["operation"] == "update":
                            updated_records |= result["record"]

                if not dry_run and self.batch_size and total_rows % self.batch_size == 0:
                    self.env.cr.commit()

            post_import_error = None
            if not blocked_error:
                try:
                    self._run_post_import_hook(
                        log=log,
                        created_records=created_records,
                        updated_records=updated_records,
                        failed_rows=failed_rows,
                    )
                except Exception as exc:
                    post_import_error = ImportRowError(str(exc), error_type="python", blocking=True)
                    self._notify_import_blockage(log=log, message=str(exc), exception=exc)

            final_error = blocked_error or post_import_error
            state = self._compute_log_state(
                success_count=success_count,
                failed_count=failed_count,
                blocked=bool(final_error),
            )
            message = self._build_summary_message(
                filename=filename,
                success_count=success_count,
                failed_count=failed_count,
                skipped_count=skipped_count,
                dry_run=dry_run,
                error=final_error,
            )
            log.sudo().write({
                "state": state,
                "total_rows": total_rows,
                "success_count": success_count,
                "failed_count": failed_count,
                "skipped_count": skipped_count,
                "finished_at": fields.Datetime.now(),
                "message": message,
            })
            runtime_vals = {"last_message": message}
            if not dry_run:
                runtime_vals["last_run_datetime"] = fields.Datetime.now()
            self.sudo().write(runtime_vals)
            return log, {
                "state": state,
                "message": message,
                "success_count": success_count,
                "failed_count": failed_count,
                "skipped_count": skipped_count,
            }
        except Exception as exc:
            message = str(exc)
            log.sudo().write({
                "state": "failed",
                "total_rows": total_rows,
                "failed_count": failed_count or total_rows,
                "finished_at": fields.Datetime.now(),
                "message": message,
            })
            self.sudo().write({"last_message": message})
            self._notify_import_blockage(log=log, message=message, exception=exc)
            return log, {
                "state": "failed",
                "message": message,
                "success_count": success_count,
                "failed_count": failed_count,
                "skipped_count": skipped_count,
            }
        finally:
            if workbook:
                workbook.close()

    def _process_single_source_row(self, source_vals, log, row_number, filename=None, dry_run=False):
        self.ensure_one()
        source_vals = dict(source_vals or {})
        hook_messages = []
        self._run_row_pre_process(source_vals, hook_messages)
        skip_row, skip_reason = self._run_row_filter(source_vals, hook_messages)
        if skip_row:
            message = skip_reason or _("Row skipped by filter.")
            self._create_log_line(
                log=log,
                row_number=row_number,
                state="skipped",
                source_vals=source_vals,
                message=message,
            )
            return {
                "status": "skipped",
                "record": self.env[self.target_model].browse(),
                "operation": "skip",
                "mapped_values": {},
            }

        mapped_vals, deferred_mappings = self._map_row_values(source_vals)
        self._validate_required_fields(source_vals, mapped_vals)
        record, operation = self._find_existing_record(source_vals, mapped_vals)

        if operation == "skip":
            message = _("Existing record found and skipped by duplicate policy.")
            self._create_log_line(
                log=log,
                row_number=row_number,
                state="skipped",
                source_vals=source_vals,
                mapped_vals=mapped_vals,
                target_record=record,
                message=message,
            )
            return {
                "status": "skipped",
                "record": record,
                "operation": operation,
                "mapped_values": mapped_vals,
            }

        self._run_row_post_process(source_vals, mapped_vals, record, hook_messages, log)

        if dry_run:
            dry_run_record = record if record and record.exists() else self.env[self.target_model].browse()
            self._create_log_line(
                log=log,
                row_number=row_number,
                state="success",
                source_vals=source_vals,
                mapped_vals=mapped_vals,
                target_record=dry_run_record,
                message=self._join_messages(hook_messages, _("Dry run completed successfully.")),
            )
            return {
                "status": "success",
                "record": dry_run_record,
                "operation": "update" if dry_run_record else "create",
                "mapped_values": mapped_vals,
            }

        target_model = self.env[self.target_model]
        try:
            if operation == "create":
                target_record = target_model.create(mapped_vals)
            else:
                record.write(mapped_vals)
                target_record = record
        except (AccessError, AccessDenied) as exc:
            raise ImportRowError(str(exc), error_type="access", mapped_vals=mapped_vals) from exc
        except ValidationError as exc:
            raise ImportRowError(str(exc), error_type="validation", mapped_vals=mapped_vals) from exc
        except UserError as exc:
            raise ImportRowError(str(exc), error_type="validation", mapped_vals=mapped_vals) from exc
        except Exception as exc:
            raise ImportRowError(str(exc), error_type="system", mapped_vals=mapped_vals) from exc

        self._apply_deferred_mappings(target_record, deferred_mappings, source_vals, mapped_vals)
        self._update_record_mapping(source_vals, mapped_vals, target_record)
        self._create_log_line(
            log=log,
            row_number=row_number,
            state="success",
            source_vals=source_vals,
            mapped_vals=mapped_vals,
            target_record=target_record,
            message=self._join_messages(hook_messages, _("Import row processed successfully.")),
        )
        return {
            "status": "success",
            "record": target_record,
            "operation": operation,
            "mapped_values": mapped_vals,
        }

    def _run_row_pre_process(self, source_vals, hook_messages):
        if not (self.pre_process_python or "").strip():
            return
        localdict = self._build_python_context(
            source_vals=source_vals,
            collector=hook_messages,
        )
        try:
            safe_eval(self.pre_process_python.strip(), localdict, mode="exec", nocopy=True)
        except Exception as exc:
            raise ImportRowError(str(exc), error_type="python") from exc

    def _run_row_filter(self, source_vals, hook_messages):
        if not (self.row_filter_python or "").strip():
            return False, False
        localdict = self._build_python_context(
            source_vals=source_vals,
            collector=hook_messages,
        )
        localdict.update({
            "skip_row": False,
            "skip_reason": False,
        })
        try:
            safe_eval(self.row_filter_python.strip(), localdict, mode="exec", nocopy=True)
        except Exception as exc:
            raise ImportRowError(str(exc), error_type="python") from exc
        return bool(localdict.get("skip_row")), localdict.get("skip_reason")

    def _run_row_post_process(self, source_vals, mapped_vals, record, hook_messages, log):
        if not (self.post_process_python or "").strip():
            return
        localdict = self._build_python_context(
            source_vals=source_vals,
            result_vals=mapped_vals,
            record=record,
            log_record=log,
            collector=hook_messages,
        )
        try:
            safe_eval(self.post_process_python.strip(), localdict, mode="exec", nocopy=True)
        except Exception as exc:
            raise ImportRowError(str(exc), error_type="python", mapped_vals=mapped_vals) from exc

    def _run_post_import_hook(self, log, created_records, updated_records, failed_rows):
        if not (self.post_import_python or "").strip():
            return
        messages = []
        localdict = self._build_python_context(
            collector=messages,
            import_log=log,
            created_records=created_records,
            updated_records=updated_records,
            failed_rows=failed_rows,
        )
        try:
            safe_eval(self.post_import_python.strip(), localdict, mode="exec", nocopy=True)
            if messages:
                log.sudo().write({
                    "message": "%s\n%s" % (log.message or "", "\n".join(messages)),
                })
        except Exception as exc:
            raise ImportRowError(str(exc), error_type="python", blocking=True) from exc

    def _map_row_values(self, source_vals):
        result_vals = {}
        deferred_mappings = []
        for mapping in self.field_mapping_ids.sorted("sequence"):
            value = self._get_source_value(source_vals, mapping)
            try:
                mapped_value = self._apply_transform(mapping, value, source_vals, result_vals)
            except ImportRowError as exc:
                exc.mapped_vals = self._json_ready(result_vals)
                raise
            except Exception as exc:
                raise ImportRowError(str(exc), error_type="validation", mapped_vals=self._json_ready(result_vals)) from exc
            if mapped_value is SKIP:
                continue
            if mapping.post_process:
                deferred_mappings.append({
                    "field_name": mapping.target_field,
                    "value": mapped_value,
                    "priority": mapping.post_process_priority,
                })
                continue
            result_vals[mapping.target_field] = mapped_value
        return result_vals, deferred_mappings

    def _get_source_value(self, source_vals, mapping):
        found = False
        value = None
        if mapping.excel_header:
            value, found = self._lookup_source_value(source_vals, mapping.excel_header)
        if not found and mapping.excel_column:
            value, found = self._lookup_source_value(source_vals, mapping.excel_column.upper())
        if not found and mapping.default_value not in (None, ""):
            value = mapping.default_value
        return value

    def _apply_transform(self, mapping, value, source_vals, result_vals):
        if mapping.transform == "skip":
            return SKIP
        if value in (None, "") and mapping.default_value not in (None, ""):
            value = mapping.default_value
        if mapping.transform == "copy":
            return value
        if mapping.transform == "char":
            return False if value in (None, "") else str(value).strip()
        if mapping.transform == "integer":
            if value in (None, ""):
                return False
            try:
                return int(float(value))
            except Exception as exc:
                raise ImportRowError(_("Invalid integer value: %s") % value, error_type="validation") from exc
        if mapping.transform in ("float", "monetary"):
            if value in (None, ""):
                return False
            try:
                return float(value)
            except Exception as exc:
                raise ImportRowError(_("Invalid numeric value: %s") % value, error_type="validation") from exc
        if mapping.transform == "boolean":
            return self._to_boolean(value, mapping)
        if mapping.transform == "date":
            return self._to_date_value(value, mapping)
        if mapping.transform == "datetime":
            return self._to_datetime_value(value, mapping)
        if mapping.transform == "selection":
            return self._to_selection_value(mapping, value)
        if mapping.transform == "many2one_name":
            return self._to_many2one_name(mapping, value)
        if mapping.transform == "many2one_ref":
            return self._to_many2one_ref(mapping, value)
        if mapping.transform == "many2many_names":
            return self._to_many2many_names(mapping, value)
        if mapping.transform == "one2many_lines":
            return self._to_one2many_lines(mapping, value)
        if mapping.transform == "python":
            if not mapping.python_expr:
                raise ImportRowError(_("Python expression is missing for mapping %s.") % mapping.target_field, error_type="blocked", blocking=True)
            localdict = self._build_python_context(
                source_vals=source_vals,
                result_vals=result_vals,
                value=value,
                mapping=mapping,
            )
            localdict["result"] = False
            try:
                safe_eval(mapping.python_expr.strip(), localdict, mode="exec", nocopy=True)
            except Exception as exc:
                raise ImportRowError(str(exc), error_type="python", mapped_vals=result_vals) from exc
            return localdict.get("result")
        raise ImportRowError(_("Unsupported transform: %s") % mapping.transform, error_type="blocked", blocking=True)

    def _validate_required_fields(self, source_vals, result_vals):
        target_model = self.env[self.target_model]
        required_field_names = [
            field_name
            for field_name, field in target_model._fields.items()
            if field_name not in TECHNICAL_FIELDS and field.required and not field.readonly
        ]
        default_values = target_model.default_get(required_field_names) if required_field_names else {}
        for mapping in self.field_mapping_ids:
            value = result_vals.get(mapping.target_field)
            if mapping.required and value in (None, False, "", [], {}):
                raise ImportRowError(
                    _("Missing required mapped field: %s") % mapping.target_field,
                    error_type="missing_required",
                    mapped_vals=result_vals,
                )
        for field_name, field in target_model._fields.items():
            if field_name in TECHNICAL_FIELDS or not field.required or field.readonly:
                continue
            if field_name not in result_vals and default_values.get(field_name) not in (None, False, "", [], {}):
                continue
            value = result_vals.get(field_name)
            if value in (None, False, "", [], {}):
                raise ImportRowError(
                    _("Required target field is missing: %s") % field_name,
                    error_type="missing_required",
                    mapped_vals=result_vals,
                )

    def _find_existing_record(self, source_vals, result_vals):
        target_model = self.env[self.target_model]
        existing_record = self._find_record_by_mapping(source_vals)
        if not existing_record and self.external_key_field:
            source_key = self._extract_external_source_key(source_vals)
            if source_key not in (None, "") and self.external_key_field in target_model._fields:
                existing_record = target_model.search([(self.external_key_field, "=", source_key)], limit=2)
                if len(existing_record) > 1:
                    raise ImportRowError(
                        _("Multiple records found for external key %s.") % source_key,
                        error_type="duplicate",
                        mapped_vals=result_vals,
                    )
        if not existing_record and result_vals.get("name") and "name" in target_model._fields:
            existing_record = target_model.search([("name", "=", result_vals["name"])], limit=2)
            if len(existing_record) > 1:
                raise ImportRowError(
                    _("Multiple records found for name %s.") % result_vals["name"],
                    error_type="duplicate",
                    mapped_vals=result_vals,
                )

        existing_record = existing_record[:1] if existing_record else target_model.browse()
        if self.import_mode == "update":
            if not existing_record:
                raise ImportRowError(
                    _("No existing record found to update."),
                    error_type="validation",
                    mapped_vals=result_vals,
                )
            return existing_record, "update"

        if existing_record:
            if self.import_mode == "create":
                if self.duplicate_policy == "error":
                    raise ImportRowError(
                        _("Duplicate record found for create-only import."),
                        error_type="duplicate",
                        mapped_vals=result_vals,
                    )
                return existing_record, "skip"
            if self.duplicate_policy == "skip":
                return existing_record, "skip"
            if self.duplicate_policy == "error":
                raise ImportRowError(
                    _("Duplicate record found and duplicate policy is set to Error."),
                    error_type="duplicate",
                    mapped_vals=result_vals,
                )
            return existing_record, "update"
        return existing_record, "create"

    def _create_failed_line(self, log, row_number, source_vals, mapped_vals=None, error_message=None, error_type="validation", filename=None):
        failed_line = self.env["ob.excel.import.failed.line"].sudo().create({
            "import_log_id": log.id,
            "row_number": row_number,
            "filename": filename or log.filename,
            "target_model": self.target_model,
            "source_values": self._json_ready(source_vals or {}),
            "mapped_values": self._json_ready(mapped_vals or {}),
            "error_type": error_type,
            "error_message": error_message or _("Unknown import error."),
        })
        return failed_line

    def _reimport_failed_line(self, failed_line):
        self.ensure_one()
        source_vals = failed_line.fixed_source_values or failed_line.source_values
        log = self._create_import_log(filename=failed_line.filename, total_rows=1)
        try:
            with self.env.cr.savepoint():
                result = self._process_single_source_row(
                    source_vals=source_vals,
                    log=log,
                    row_number=failed_line.row_number,
                    filename=failed_line.filename,
                    dry_run=False,
                )
            failed_line.sudo().write({
                "state": "reimported",
                "target_record_id": result["record"].id if result["record"] else False,
                "reimported_at": fields.Datetime.now(),
            })
            log.sudo().write({
                "state": "done",
                "success_count": 1,
                "finished_at": fields.Datetime.now(),
                "message": _("Failed line re-imported successfully."),
            })
        except Exception as exc:
            error = self._normalize_row_exception(exc)
            failed_line.sudo().write({
                "state": "failed_again",
                "error_message": error.message,
            })
            log.sudo().write({
                "state": "failed",
                "failed_count": 1,
                "finished_at": fields.Datetime.now(),
                "message": error.message,
            })
            self._notify_import_blockage(
                log=log,
                message=error.message,
                failed_line=failed_line,
                exception=exc,
            )

    def _reimport_failed_lines(self, failed_lines):
        for line in failed_lines:
            self._reimport_failed_line(line)

    def _notify_import_blockage(self, log, message, failed_line=None, exception=None):
        self.ensure_one()
        if not self.notify_on_error and not self.create_activity_on_error:
            return
        summary = _("Excel Import Error: %s") % self.name
        note = _(
            "An error occurred during Excel import.<br/>"
            "<b>Template:</b> %s<br/>"
            "<b>File:</b> %s<br/>"
            "<b>Message:</b> %s"
        ) % (
            self.name,
            log.filename or "",
            message or "",
        )
        if failed_line:
            note += _("<br/><b>Row:</b> %s") % failed_line.row_number
        if exception:
            _logger.exception("Excel import blockage on template %s: %s", self.display_name, exception)
        if self.notify_on_error and self.responsible_user_id.partner_id:
            self.sudo().message_post(
                body=note,
                subject=summary,
                partner_ids=[self.responsible_user_id.partner_id.id],
                message_type="notification",
                subtype_xmlid="mail.mt_comment",
            )
        if self.create_activity_on_error and self.activity_type_id and self.responsible_user_id:
            self.sudo().activity_schedule(
                activity_type_id=self.activity_type_id.id,
                user_id=self.responsible_user_id.id,
                summary=summary,
                note=note,
            )
        if failed_line and not failed_line.activity_created:
            failed_line.sudo().activity_created = True

    def action_import_from_ftp(self):
        self.ensure_one()
        summary = self._run_ftp_import_core()
        notif_type = "success"
        if summary["failed_files"]:
            notif_type = "warning" if summary["imported_files"] else "danger"
        return self._build_notification_action(
            title=_("FTP Import"),
            message=summary["message"],
            notif_type=notif_type,
        )

    def _run_ftp_import_core(self):
        self.ensure_one()
        try:
            self._validate_ftp_configuration()
        except Exception as exc:
            message = str(exc)
            log = self._create_system_failed_log(message)
            self._notify_import_blockage(log=log, message=message, exception=exc)
            return {
                "processed_files": 0,
                "imported_files": 0,
                "failed_files": 1,
                "message": message,
            }

        processed_files = 0
        imported_files = 0
        failed_files = 0
        connection = None
        try:
            connection = self._ftp_connect()
            files = self._fetch_ftp_files(connection)
            if not files:
                message = _("No matching files were found on the remote path.")
                self.sudo().write({"last_message": message})
                return {
                    "processed_files": 0,
                    "imported_files": 0,
                    "failed_files": 0,
                    "message": message,
                }
            for file_info in files:
                processed_files += 1
                try:
                    log, summary = self._run_import_core(
                        file_content=file_info["content"],
                        filename=file_info["filename"],
                        dry_run=False,
                    )
                    if summary["state"] in ("done", "partial"):
                        imported_files += 1
                        self._archive_or_delete_ftp_file(connection, file_info["filename"])
                    else:
                        failed_files += 1
                        self._notify_import_blockage(log=log, message=summary["message"])
                except Exception as exc:
                    failed_files += 1
                    message = str(exc)
                    log = self._create_system_failed_log(message, filename=file_info["filename"])
                    self._notify_import_blockage(log=log, message=message, exception=exc)
            message = _("Processed %s file(s): %s imported, %s failed.") % (
                processed_files,
                imported_files,
                failed_files,
            )
            self.sudo().write({
                "last_message": message,
                "last_run_datetime": fields.Datetime.now(),
            })
            return {
                "processed_files": processed_files,
                "imported_files": imported_files,
                "failed_files": failed_files,
                "message": message,
            }
        except Exception as exc:
            message = str(exc)
            log = self._create_system_failed_log(message)
            self._notify_import_blockage(log=log, message=message, exception=exc)
            return {
                "processed_files": processed_files,
                "imported_files": imported_files,
                "failed_files": failed_files + 1,
                "message": message,
            }
        finally:
            self._close_ftp_connection(connection)

    def _ftp_connect(self):
        self.ensure_one()
        if self.ftp_protocol == "ftp":
            connection = ftplib.FTP()
            connection.connect(self.ftp_host, self.ftp_port or 21, timeout=30)
            connection.login(self.ftp_username or "", self.ftp_password or "")
            connection.set_pasv(self.ftp_passive_mode)
            connection.cwd(self.ftp_path or "/")
            return {
                "protocol": "ftp",
                "client": connection,
            }
        if not paramiko:
            raise UserError(_("The Python package 'paramiko' is required for SFTP imports."))
        transport = paramiko.Transport((self.ftp_host, self.ftp_port or 22))
        transport.connect(username=self.ftp_username or "", password=self.ftp_password or "")
        client = paramiko.SFTPClient.from_transport(transport)
        client.chdir(self.ftp_path or ".")
        return {
            "protocol": "sftp",
            "client": client,
            "transport": transport,
        }

    def _fetch_ftp_files(self, connection):
        self.ensure_one()
        files = []
        pattern = self.ftp_filename_pattern or "*.xlsx"
        protocol = connection["protocol"]
        client = connection["client"]
        if protocol == "ftp":
            names = sorted(client.nlst())
            for filename in names:
                if not fnmatch.fnmatch(filename, pattern):
                    continue
                buffer = io.BytesIO()
                client.retrbinary("RETR %s" % filename, buffer.write)
                files.append({
                    "filename": filename,
                    "content": buffer.getvalue(),
                })
        else:
            for filename in sorted(client.listdir(".")):
                if not fnmatch.fnmatch(filename, pattern):
                    continue
                with client.open(filename, "rb") as file_handle:
                    files.append({
                        "filename": filename,
                        "content": file_handle.read(),
                    })
        return files

    def _archive_or_delete_ftp_file(self, connection, filename):
        protocol = connection["protocol"]
        client = connection["client"]
        if self.ftp_archive_path:
            archive_name = "%s_%s" % (fields.Datetime.now().strftime("%Y%m%d_%H%M%S"), filename)
            archive_dir = self.ftp_archive_path
            self._ensure_remote_directory(connection, archive_dir)
            destination = posixpath.join(archive_dir, archive_name)
            if protocol == "ftp":
                client.rename(filename, destination)
            else:
                client.rename(filename, destination)
            return
        if not self.ftp_delete_after_import:
            return
        if protocol == "ftp":
            client.delete(filename)
        else:
            client.remove(filename)

    def _ensure_remote_directory(self, connection, directory):
        if not directory:
            return
        protocol = connection["protocol"]
        client = connection["client"]
        current = ""
        for part in [segment for segment in directory.split("/") if segment]:
            current = "%s/%s" % (current, part) if current else part
            try:
                if protocol == "ftp":
                    client.mkd(current)
                else:
                    client.stat(current)
            except Exception:
                if protocol == "ftp":
                    try:
                        client.mkd(current)
                    except Exception:
                        pass
                else:
                    try:
                        client.mkdir(current)
                    except Exception:
                        pass

    @api.model
    def _cron_run_due_imports(self):
        now = fields.Datetime.now()
        templates = self.search([
            ("active", "=", True),
            ("auto_import", "=", True),
            ("ftp_enabled", "=", True),
            ("next_execution_datetime", "<=", now),
        ])
        for template in templates:
            try:
                template._run_ftp_import_core()
                template.sudo().last_run_datetime = fields.Datetime.now()
                template.sudo()._update_next_execution()
            except Exception as exc:
                log = template._create_system_failed_log(str(exc))
                template._notify_import_blockage(
                    log=log,
                    message=str(exc),
                    exception=exc,
                )

    def _get_interval_delta(self):
        self.ensure_one()
        interval = self.interval_number or 1
        if self.interval_unit == "minute":
            return timedelta(minutes=interval)
        if self.interval_unit == "hour":
            return timedelta(hours=interval)
        if self.interval_unit == "day":
            return timedelta(days=interval)
        if self.interval_unit == "week":
            return timedelta(weeks=interval)
        if self.interval_unit == "month":
            return relativedelta(months=interval)
        return timedelta(days=1)

    def _update_next_execution(self, from_dt=None):
        for template in self:
            if template.auto_import and template.ftp_enabled:
                base_dt = from_dt or fields.Datetime.now()
                template.next_execution_datetime = base_dt + template._get_interval_delta()
            else:
                template.next_execution_datetime = False

    def _read_excel_file(self, binary_file, filename):
        if not filename or not filename.lower().endswith(".xlsx"):
            raise UserError(_("Only .xlsx files are supported."))
        try:
            return load_workbook(filename=io.BytesIO(binary_file), read_only=True, data_only=True)
        except Exception as exc:
            raise UserError(_("Unable to read Excel file: %s") % exc) from exc

    def _get_workbook_sheet(self, workbook):
        self.ensure_one()
        if self.sheet_name:
            if self.sheet_name not in workbook.sheetnames:
                raise UserError(_("Sheet '%s' was not found in the workbook.") % self.sheet_name)
            return workbook[self.sheet_name]
        return workbook.active

    def _extract_headers(self, sheet):
        if self.header_row < 1:
            raise UserError(_("Header row must be greater than zero."))
        headers = []
        for cell in sheet[self.header_row]:
            value = self._normalize_cell_value(cell.value)
            headers.append(str(value).strip() if value not in (None, "") else False)
        return headers

    def _iter_excel_rows(self, sheet, headers):
        for row_index in range(self.data_start_row, sheet.max_row + 1):
            row_vals = {"_row_number": row_index}
            has_value = False
            row = sheet[row_index]
            for column_index, cell in enumerate(row, start=1):
                normalized_value = self._normalize_cell_value(cell.value)
                if normalized_value not in (None, ""):
                    has_value = True
                column_letter = get_column_letter(column_index)
                row_vals[column_letter] = normalized_value
                header_name = headers[column_index - 1] if column_index - 1 < len(headers) else False
                if header_name:
                    row_vals[header_name] = normalized_value
            if has_value:
                yield row_vals

    def _normalize_cell_value(self, value):
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return value
        if isinstance(value, time):
            return value
        if isinstance(value, str):
            return value.strip()
        return value

    def _validate_template_configuration(self):
        self.ensure_one()
        self._validate_template_model()
        if self.header_row < 1 or self.data_start_row < 1:
            raise UserError(_("Header row and data start row must be greater than zero."))
        if self.data_start_row <= self.header_row:
            raise UserError(_("Data start row must be greater than header row."))
        if not self.field_mapping_ids:
            raise UserError(_("Please configure at least one field mapping."))
        for mapping in self.field_mapping_ids:
            if not (mapping.excel_header or mapping.excel_column or mapping.default_value not in (None, "")):
                raise UserError(_("Each mapping must define an Excel header, Excel column, or a default value."))
            if mapping.target_field not in self.env[self.target_model]._fields:
                raise UserError(_("Target field '%s' does not exist on model %s.") % (mapping.target_field, self.target_model))

    def _validate_template_model(self):
        if not self.target_model or not self.env.registry.get(self.target_model):
            raise UserError(_("Target model is not available."))

    def _validate_ftp_configuration(self):
        self.ensure_one()
        required_values = {
            _("FTP host"): self.ftp_host,
            _("FTP username"): self.ftp_username,
        }
        for label, value in required_values.items():
            if not value:
                raise UserError(_("%s is required for FTP imports.") % label)
        if self.ftp_protocol == "sftp" and not paramiko:
            raise UserError(_("The Python package 'paramiko' is required for SFTP imports."))

    def _decode_binary_file(self, file_content):
        if not file_content:
            raise UserError(_("An Excel file is required."))
        if isinstance(file_content, bytes):
            try:
                return base64.b64decode(file_content, validate=True)
            except Exception:
                return file_content
        if isinstance(file_content, str):
            return base64.b64decode(file_content)
        raise UserError(_("Unsupported file content format."))

    def _build_python_context(self, **extra):
        collector = extra.pop("collector", None)

        def hook_log(message, level="info"):
            if collector is not None:
                collector.append("[%s] %s" % (level, message))
            log_method = getattr(_logger, level, _logger.info)
            log_method("%s", message)

        context = {
            "env": self.env,
            "template": self,
            "datetime": datetime,
            "date": date,
            "time": time,
            "relativedelta": relativedelta,
            "UserError": UserError,
            "Command": Command,
            "log": hook_log,
            "_logger": _logger,
        }
        context.update(extra)
        return context

    def _to_boolean(self, value, mapping):
        if value in (None, ""):
            return False
        normalized = str(value).strip().lower()
        true_values = {item.strip().lower() for item in (mapping.boolean_true_values or "").split(",") if item.strip()}
        false_values = {item.strip().lower() for item in (mapping.boolean_false_values or "").split(",") if item.strip()}
        if normalized in true_values:
            return True
        if normalized in false_values:
            return False
        raise ImportRowError(_("Invalid boolean value: %s") % value, error_type="validation")

    def _to_date_value(self, value, mapping):
        if value in (None, ""):
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            if mapping.date_format:
                return datetime.strptime(str(value), mapping.date_format).date()
            return date_parser.parse(str(value)).date()
        except Exception as exc:
            raise ImportRowError(_("Invalid date value: %s") % value, error_type="validation") from exc

    def _to_datetime_value(self, value, mapping):
        if value in (None, ""):
            return False
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, time.min)
        try:
            if mapping.date_format:
                return datetime.strptime(str(value), mapping.date_format)
            return date_parser.parse(str(value))
        except Exception as exc:
            raise ImportRowError(_("Invalid datetime value: %s") % value, error_type="validation") from exc

    def _to_selection_value(self, mapping, value):
        if value in (None, ""):
            return False
        field = self.env[self.target_model]._fields[mapping.target_field]
        selection_values = field._description_selection(self.env)
        normalized = str(value).strip().lower()
        for key, label in selection_values:
            if normalized in {str(key).lower(), str(label).lower()}:
                return key
        raise ImportRowError(
            _("Invalid selection value '%s' for field %s.") % (value, mapping.target_field),
            error_type="validation",
        )

    def _to_many2one_name(self, mapping, value):
        if value in (None, ""):
            return False
        field = self.env[self.target_model]._fields[mapping.target_field]
        comodel = self.env[field.comodel_name]
        search_field = mapping.m2o_search_field or "name"
        records = comodel.search([(search_field, "=", value)], limit=2)
        if len(records) > 1:
            raise ImportRowError(
                _("Multiple related records found for %s.") % value,
                error_type="missing_relation",
            )
        if records:
            return records.id
        if mapping.missing_m2o_policy == "skip":
            return SKIP
        if mapping.missing_m2o_policy == "create":
            return comodel.create({search_field: value}).id
        raise ImportRowError(
            _("Related record not found for %s.") % value,
            error_type="missing_relation",
        )

    def _to_many2one_ref(self, mapping, value):
        if value in (None, ""):
            return False
        field = self.env[self.target_model]._fields[mapping.target_field]
        comodel_name = field.comodel_name
        comodel = self.env[comodel_name]
        string_value = str(value).strip()
        if "." in string_value:
            try:
                record = self.env.ref(string_value, raise_if_not_found=False)
                if record and record._name == comodel_name:
                    return record.id
            except Exception:
                pass
        record_map = self.env["ob.excel.import.record.map"].sudo().search([
            ("template_id", "=", self.id),
            ("source_key", "=", string_value),
            ("target_model", "=", comodel_name),
            ("active", "=", True),
        ], limit=1)
        if record_map:
            record = comodel.browse(record_map.target_record_id).exists()
            if record:
                return record.id
        candidate_fields = ["barcode", "code", mapping.m2o_search_field or "name"]
        for field_name in candidate_fields:
            if field_name not in comodel._fields:
                continue
            record = comodel.search([(field_name, "=", string_value)], limit=2)
            if len(record) > 1:
                raise ImportRowError(
                    _("Multiple related records found for %s.") % string_value,
                    error_type="missing_relation",
                )
            if record:
                return record.id
        if mapping.missing_m2o_policy == "skip":
            return SKIP
        if mapping.missing_m2o_policy == "create":
            search_field = mapping.m2o_search_field or "name"
            if search_field not in comodel._fields:
                search_field = "name"
            return comodel.create({search_field: string_value}).id
        raise ImportRowError(
            _("Related reference not found for %s.") % string_value,
            error_type="missing_relation",
        )

    def _to_many2many_names(self, mapping, value):
        if value in (None, ""):
            return [(6, 0, [])]
        field = self.env[self.target_model]._fields[mapping.target_field]
        comodel = self.env[field.comodel_name]
        search_field = mapping.m2o_search_field or "name"
        tokens = value if isinstance(value, (list, tuple)) else str(value).split(mapping.x2many_separator or ",")
        ids = []
        for token in tokens:
            token = token.strip() if isinstance(token, str) else token
            if token in (None, ""):
                continue
            record = comodel.search([(search_field, "=", token)], limit=2)
            if len(record) > 1:
                raise ImportRowError(
                    _("Multiple related records found for %s.") % token,
                    error_type="missing_relation",
                )
            if not record:
                if mapping.missing_m2o_policy == "skip":
                    continue
                if mapping.missing_m2o_policy == "create":
                    record = comodel.create({search_field: token})
                else:
                    raise ImportRowError(
                        _("Related record not found for %s.") % token,
                        error_type="missing_relation",
                    )
            ids.append(record.id)
        return [(6, 0, ids)]

    def _to_one2many_lines(self, mapping, value):
        if value in (None, ""):
            return []
        if isinstance(value, str):
            try:
                value = json.loads(value)
            except json.JSONDecodeError as exc:
                raise ImportRowError(_("One2many value must be valid JSON."), error_type="validation") from exc
        if isinstance(value, dict):
            value = [value]
        if not isinstance(value, list):
            raise ImportRowError(_("One2many value must be a list or dictionary."), error_type="validation")
        commands = []
        for line_vals in value:
            if isinstance(line_vals, dict):
                commands.append(Command.create(line_vals))
            elif isinstance(line_vals, (list, tuple)) and len(line_vals) == 3:
                commands.append(tuple(line_vals))
            else:
                raise ImportRowError(_("Invalid one2many line definition."), error_type="validation")
        return commands

    def _apply_deferred_mappings(self, record, deferred_mappings, source_vals, mapped_vals):
        for deferred in sorted(deferred_mappings, key=lambda item: item["priority"]):
            try:
                record.write({deferred["field_name"]: deferred["value"]})
            except Exception as exc:
                raise ImportRowError(
                    _("Deferred field update failed for %s: %s") % (deferred["field_name"], exc),
                    error_type="validation",
                    mapped_vals=mapped_vals,
                ) from exc

    def _update_record_mapping(self, source_vals, mapped_vals, record):
        source_key = self._extract_external_source_key(source_vals)
        if source_key in (None, ""):
            return
        map_model = self.env["ob.excel.import.record.map"].sudo()
        existing_map = map_model.search([
            ("template_id", "=", self.id),
            ("source_key", "=", str(source_key)),
            ("target_model", "=", self.target_model),
        ], limit=1)
        values = {
            "template_id": self.id,
            "source_key": str(source_key),
            "target_model": self.target_model,
            "target_record_id": record.id,
            "active": True,
        }
        if existing_map:
            existing_map.write(values)
        else:
            map_model.create(values)

    def _find_record_by_mapping(self, source_vals):
        source_key = self._extract_external_source_key(source_vals)
        if source_key in (None, ""):
            return self.env[self.target_model].browse()
        record_map = self.env["ob.excel.import.record.map"].sudo().search([
            ("template_id", "=", self.id),
            ("source_key", "=", str(source_key)),
            ("target_model", "=", self.target_model),
            ("active", "=", True),
        ], limit=1)
        if not record_map:
            return self.env[self.target_model].browse()
        return self.env[self.target_model].browse(record_map.target_record_id).exists()

    def _extract_external_source_key(self, source_vals):
        if not self.external_key_column:
            return False
        value, found = self._lookup_source_value(source_vals, self.external_key_column)
        return value if found else False

    def _lookup_source_value(self, source_vals, key):
        if key in source_vals:
            return source_vals[key], True
        normalized = str(key).strip().lower()
        for existing_key, value in (source_vals or {}).items():
            if not isinstance(existing_key, str) or existing_key.startswith("_"):
                continue
            if existing_key.strip().lower() == normalized:
                return value, True
        return False, False

    def _create_import_log(self, filename, total_rows=0):
        return self.env["ob.excel.import.log"].sudo().create({
            "template_id": self.id,
            "filename": filename,
            "state": "running",
            "started_at": fields.Datetime.now(),
            "total_rows": total_rows,
        })

    def _create_system_failed_log(self, message, filename=False):
        return self.env["ob.excel.import.log"].sudo().create({
            "template_id": self.id,
            "filename": filename,
            "state": "failed",
            "started_at": fields.Datetime.now(),
            "finished_at": fields.Datetime.now(),
            "message": message,
        })

    def _create_log_line(self, log, row_number, state, source_vals, mapped_vals=None, target_record=None, message=None):
        return self.env["ob.excel.import.log.line"].sudo().create({
            "log_id": log.id,
            "row_number": row_number,
            "state": state,
            "target_model": self.target_model,
            "target_record_id": target_record.id if target_record else False,
            "source_values": self._json_ready(source_vals or {}),
            "mapped_values": self._json_ready(mapped_vals or {}),
            "message": message,
        })

    def _normalize_row_exception(self, exc):
        if isinstance(exc, ImportRowError):
            return exc
        if isinstance(exc, (AccessError, AccessDenied)):
            return ImportRowError(str(exc), error_type="access")
        if isinstance(exc, (UserError, ValidationError)):
            return ImportRowError(str(exc), error_type="validation")
        return ImportRowError(str(exc), error_type="system")

    def _json_ready(self, value):
        if isinstance(value, models.BaseModel):
            return value.ids
        if isinstance(value, dict):
            return {str(key): self._json_ready(val) for key, val in value.items()}
        if isinstance(value, (list, tuple)):
            return [self._json_ready(item) for item in value]
        if isinstance(value, datetime):
            return value.isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, time):
            return value.isoformat()
        if isinstance(value, bytes):
            return base64.b64encode(value).decode()
        return value

    def _compute_log_state(self, success_count, failed_count, blocked=False):
        if blocked and not success_count:
            return "failed"
        if failed_count and success_count:
            return "partial"
        if failed_count:
            return "failed"
        return "done"

    def _build_summary_message(self, filename, success_count, failed_count, skipped_count, dry_run=False, error=None):
        action_label = _("Dry run") if dry_run else _("Import")
        message = _(
            "%(action)s finished for %(filename)s. Success: %(success)s, Failed: %(failed)s, Skipped: %(skipped)s."
        ) % {
            "action": action_label,
            "filename": filename,
            "success": success_count,
            "failed": failed_count,
            "skipped": skipped_count,
        }
        if error:
            message = "%s %s" % (message, error.message)
        return message

    def _build_notification_action(self, title, message, notif_type="success", next_action=False):
        params = {
            "title": title,
            "message": message,
            "type": notif_type,
            "sticky": notif_type in ("danger", "warning"),
        }
        if next_action:
            params["next"] = next_action
        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": params,
        }

    def _guess_transform_for_field(self, field):
        transform_map = {
            "char": "char",
            "text": "char",
            "html": "char",
            "integer": "integer",
            "float": "float",
            "monetary": "monetary",
            "boolean": "boolean",
            "date": "date",
            "datetime": "datetime",
            "selection": "selection",
            "many2one": "many2one_name",
            "many2many": "many2many_names",
            "one2many": "one2many_lines",
        }
        return transform_map.get(field.type, "copy")

    def _join_messages(self, messages, base_message):
        if not messages:
            return base_message
        return "%s\n%s" % (base_message, "\n".join(messages))

    def _close_ftp_connection(self, connection):
        if not connection:
            return
        try:
            if connection["protocol"] == "ftp":
                connection["client"].quit()
            else:
                connection["client"].close()
                connection["transport"].close()
        except Exception:
            pass
